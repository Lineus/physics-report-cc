#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""逐页校对：docx + PDF 一起查，并把每页渲成图供人/多模态模型看。

用法：
    python3 proof.py "<目录>" "*第5周.docx" [--dpi 110] [--render <输出目录>]

查六件事：
  1. 空白页      每页文字量 < 50 字 → 疑似空白
  2. 图溢出      图片 bbox 越过页边距（20pt）或超出页面
  3. 图注落单    图注所在页没有任何图片 → 图注跑到次页
  4. 内容缺失    docx 里的正文段在 PDF 文本里找不到
  5. 字数        综述 / 重点 / 难点按锚点切段实测
  6. 公式        docx 里含 OMML 的段数（配合看图确认没变方框）

退出码：0 无 WARN；1 有 WARN。
"""
import argparse
import glob
import os
import re
import sys

import fitz

W_T = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t"
M_T = "{http://schemas.openxmlformats.org/officeDocument/2006/math}t"
CAP_RE = re.compile(r"^图\s*\d")
MARGIN_PT = 20
MIN_CHARS = 50


def norm(s):
    return re.sub(r"\s+", "", s)


def para_text(p):
    return "".join(n.text for n in p._p.iter()
                   if n.tag in (W_T, M_T) and n.text).strip()


def lead_text(p):
    """公式之前的中文部分。

    PDF 提取出的公式字符与 docx 里的 OMML 拼接方式不同（多空格、换行、
    上下标错位），整段比对必然误报。只用「第一个公式之前」的文字做核对。
    """
    buf = []
    for node in p._p.iter():
        if node.tag == M_T:
            break
        if node.tag == W_T and node.text:
            buf.append(node.text)
    return "".join(buf).strip()


def docx_info(path):
    import docx
    d = docx.Document(path)
    paras = [para_text(p) for p in d.paragraphs]
    caps, seen_img = [], False
    for p, t in zip(d.paragraphs, paras):
        if "graphic" in p._p.xml:
            seen_img = True
        elif seen_img and CAP_RE.match(t):
            caps.append(t)
            seen_img = False
    math_paras = sum(1 for p in d.paragraphs
                     if any(n.tag == M_T for n in p._p.iter()))
    leads = [lead_text(p) for p in d.paragraphs]
    return d, paras, caps, math_paras, leads


def span(paras, start, stop):
    try:
        s = next(i for i, t in enumerate(paras) if t.startswith(start))
        e = next(i for i, t in enumerate(paras) if t.startswith(stop))
    except StopIteration:
        return None
    return sum(len(norm(t)) for t in paras[s + 1:e])


def check(docx_path, pdf_path, dpi, render_dir):
    d, paras, caps, math_paras, leads = docx_info(docx_path)
    pdf = fitz.open(pdf_path)
    warns = 0

    print("文件：%s" % os.path.basename(docx_path))
    print("页数：%d   docx 正文段：%d   图注：%d   含公式段：%d"
          % (len(pdf), len([t for t in paras if t]), len(caps), math_paras))
    print()

    full = ""
    for i, page in enumerate(pdf, 1):
        txt = page.get_text()
        full += txt
        chars = len(norm(txt))
        imgs = page.get_image_info()
        flags = []

        if chars < MIN_CHARS and not imgs:
            flags.append("疑似空白页（%d 字、无图）" % chars)
        elif chars < MIN_CHARS:
            flags.append("文字极少（%d 字）" % chars)

        pw, ph = page.rect.width, page.rect.height
        for im in imgs:
            x0, y0, x1, y1 = im["bbox"]
            if x0 < MARGIN_PT or y0 < MARGIN_PT or x1 > pw - MARGIN_PT or y1 > ph - MARGIN_PT:
                flags.append("图片贴边/溢出（%.0f,%.0f,%.0f,%.0f）" % (x0, y0, x1, y1))
                break

        page_caps = [c for c in caps if norm(c)[:10] and norm(c)[:10] in norm(txt)]
        if page_caps and not imgs:
            flags.append("图注在本页但本页无图：%s" % page_caps[0][:20])

        if render_dir:
            os.makedirs(render_dir, exist_ok=True)
            page.get_pixmap(dpi=dpi).save(
                os.path.join(render_dir, "p%02d.png" % i))

        mark = "WARN" if flags else "OK  "
        if flags:
            warns += len(flags)
        print("%s p%-2d %4d 字  %d 图  %s" % (mark, i, chars, len(imgs), "；".join(flags)))

    print()
    ptxt = norm(full)
    missing = []
    for t, lead in zip(paras, leads):
        key = norm(lead) if len(norm(lead)) >= 8 else ""
        if not key:
            continue                      # 纯公式段，或公式开头，跳过
        probe = key[:20]
        if probe not in ptxt:
            missing.append(lead)
    if missing:
        warns += len(missing)
        print("WARN docx 有 %d 段在 PDF 里找不到（可能被截断或分页拆开）：" % len(missing))
        for t in missing[:5]:
            print("     " + t[:60])

    for label, a, b in [("综述", "（自述实验现象", "2.实验重点"),
                        ("重点", "（简述本实验的学习重点", "3.实验难点"),
                        ("难点", "（简述本实验的实现难点", "注意事项")]:
        n = span(paras, a, b)
        if n is not None:
            print("%s = %d 字" % (label, n))

    if render_dir:
        print("\n页面图：%s（逐张看：溢出 / 断图 / 图注跑页首 / 整页空白 / 公式方框）" % render_dir)
    print("\n结论：%s" % ("无异常" if not warns else "%d 处需人工确认" % warns))
    return 0 if not warns else 1


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("directory")
    ap.add_argument("pattern", nargs="?", default="*.docx")
    ap.add_argument("--dpi", type=int, default=110)
    ap.add_argument("--render", default="")
    a = ap.parse_args()

    base = a.directory
    m = re.match(r"^([A-Za-z]):[\\/](.*)$", base)
    if m and os.path.isdir("/mnt"):
        base = "/mnt/%s/%s" % (m.group(1).lower(), m.group(2).replace("\\", "/"))

    docx = [p for p in glob.glob(os.path.join(base, a.pattern))
            if not os.path.basename(p).startswith("~$")]
    if not docx:
        sys.exit("没找到 docx：%s/%s" % (base, a.pattern))
    docx.sort(key=os.path.getmtime, reverse=True)
    src = docx[0]
    pdf = os.path.splitext(src)[0] + ".pdf"
    if not os.path.exists(pdf):
        sys.exit("没有对应的 PDF，先跑 export_pdf.py")
    return check(src, pdf, a.dpi, a.render)


if __name__ == "__main__":
    sys.exit(main())
