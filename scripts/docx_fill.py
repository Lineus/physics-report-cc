# -*- coding: utf-8 -*-
"""按锚点填 Word 模板（python-docx）：中文样式、真上下标、插图。

依赖：python-docx。用法见 docx-report-build/SKILL.md Step 3。
"""
import re

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

BODY_FONT = "宋体"


# ---------------------------------------------------------------- 查看/定位
def outline(doc):
    """打印每段的索引 / 文本 / 样式——先看清结构再动。"""
    for i, p in enumerate(doc.paragraphs):
        print(i, repr(p.text[:40]), p.style.name, "img" if "graphic" in p._p.xml else "")


def find(doc, prefix):
    """按开头文字找锚点段。找不到就报错，别猜索引。"""
    for p in doc.paragraphs:
        if p.text.strip().startswith(prefix):
            return p
    raise SystemExit(f"锚点未找到: {prefix}")


def clear_between(start, stop):
    """删掉两个锚点段之间的占位段（模板的空白行）。"""
    body = start._p.getparent()
    els = list(body)
    for el in els[els.index(start._p) + 1:els.index(stop._p)]:
        if el.tag.endswith("}p"):
            body.remove(el)


# ---------------------------------------------------------------- 写入
def style_run(r, size=12, bold=False, font=BODY_FONT):
    """中文字体必须连 eastAsia 一起设，否则会掉回默认西文字体。"""
    r.font.size = Pt(size)
    r.bold = bold
    r.font.name = font
    r._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    return r


def add_rich(p, text, size=12, bold=False):
    """mini-markup：U_{p-p} → 下标，e^{-t/RC} → 上标。"""
    for tk in re.split(r"(\^\{[^}]*\}|_\{[^}]*\})", text):
        if not tk:
            continue
        if tk.startswith("_{"):
            style_run(p.add_run(tk[2:-1]), size, bold).font.subscript = True
        elif tk.startswith("^{"):
            style_run(p.add_run(tk[2:-1]), size, bold).font.superscript = True
        else:
            style_run(p.add_run(tk), size, bold)


def new_para(doc, anchor):
    """建一个空段并插到 anchor 之后（先建在文末再搬，顺序插入即正序）。"""
    p = doc.add_paragraph()
    anchor._p.addnext(p._p)
    return p


def text_block(doc, anchor, text, bold=False, indent=True, size=12, line_spacing=1.25):
    p = new_para(doc, anchor)
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = line_spacing
    if indent:
        pf.first_line_indent = Pt(2 * size)      # 首行缩进 2 字符
    add_rich(p, text, size=size, bold=bold)
    return p


def figure_block(doc, anchor, path, width_cm, caption, size=10.5, gap_pt=6):
    """居中插图 + 紧跟一条居中图注；返回图注段（作为下一个插入锚点）。

    只给图段设 keep_with_next：图与图注不拆页；图注段不设，否则
    图→注→下一张图… 会串成一条链，整组图被挤到一页。
    """
    p = new_para(doc, anchor)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(gap_pt)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(path, width=Cm(width_cm))
    c = new_para(doc, p)
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.paragraph_format.space_before = Pt(0)
    c.paragraph_format.space_after = Pt(gap_pt)
    add_rich(c, caption, size=size)
    return c


# ---------------------------------------------------------------- 核对
def count_chars(text):
    """剥掉 mini-markup 后数字数（模板里"不超过 N 字"用这个核）。"""
    return len(re.sub(r"[\^_]\{([^}]*)\}", r"\1", text))


def text_width_cm(doc):
    """版心宽度，用来判断插图宽上限（A4 本机约 17.2cm）。"""
    s = doc.sections[0]
    return (s.page_width - s.left_margin - s.right_margin) / 360000


def open_base(template, target, fallback=None):
    """从模板起一份新稿；模板不在了就退回上一版成品当底稿。

    源目录可能同时有别的 agent 或他自己在动文件（模板被移走、成品被改名都
    遇到过）；重建时 clear_between() 会清掉旧内容，所以拿上一版成品当底稿是
    幂等的。改文字只改脚本里的文本块，别手改 docx。
    """
    import os
    import shutil

    base = template if os.path.exists(template) else fallback
    if not base or not os.path.exists(base):
        raise SystemExit(f"底稿不存在：{template} / {fallback}")
    if os.path.abspath(base) != os.path.abspath(target):
        shutil.copyfile(base, target)
    return docx.Document(target)
